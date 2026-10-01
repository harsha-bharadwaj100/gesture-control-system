import docx
import pypdf
import os

def dump_docx(path, out_file):
    try:
        doc = docx.Document(path)
        fullText = []
        for para in doc.paragraphs:
            if para.text.strip():
                fullText.append(para.text)
        for t_idx, table in enumerate(doc.tables):
            fullText.append(f"\n--- Table {t_idx+1} ---")
            for row in table.rows:
                fullText.append(' | '.join([cell.text.strip().replace('\n', ' ') for cell in row.cells]))
        with open(out_file, 'w', encoding='utf-8') as f:
            f.write('\n'.join(fullText))
        print(f"Dumped {path} -> {out_file}")
    except Exception as e:
        print(f"Error dumping {path}: {e}")

def dump_pdf(path, out_file):
    try:
        reader = pypdf.PdfReader(path)
        text = ''
        for i, page in enumerate(reader.pages):
            text += f'\n=== Page {i+1} ===\n' + page.extract_text()
        with open(out_file, 'w', encoding='utf-8') as f:
            f.write(text)
        print(f"Dumped {path} -> {out_file}")
    except Exception as e:
        print(f"Error dumping {path}: {e}")

dump_docx('benchmark_analysis_report.docx', 'dump_benchmark_report.txt')
dump_docx('EMG_Research_Documentation_Checklist.docx', 'dump_checklist.txt')

for i in range(1, 6):
    fname = f'Gesture_Paper ({i}).pdf'
    if os.path.exists(fname):
        dump_pdf(fname, f'dump_gesture_paper_{i}.txt')

if os.path.exists('Gesture_Paper.pdf'):
    dump_pdf('Gesture_Paper.pdf', 'dump_gesture_paper_0.txt')
