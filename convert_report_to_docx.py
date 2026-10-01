import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def convert_md_to_docx(md_path="Results/benchmark_analysis_report.md", docx_path="Results/benchmark_analysis_report.docx"):
    if not os.path.exists(md_path):
        print(f"Error: {md_path} not found.")
        return

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = docx.Document()
    
    # Page Setup - Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles Setup
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    in_table = False
    table_headers = []
    table_rows = []

    def flush_table():
        nonlocal in_table, table_headers, table_rows
        if not table_headers:
            return
        
        num_cols = len(table_headers)
        table = doc.add_table(rows=len(table_rows) + 1, cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # Format Header Row
        hdr_cells = table.rows[0].cells
        for i, header_text in enumerate(table_headers):
            hdr_cells[i].text = header_text.strip('*').strip()
            set_cell_background(hdr_cells[i], "1F4E78") # Professional Navy Blue Header
            set_cell_margins(hdr_cells[i], top=120, bottom=120, left=150, right=150)
            p = hdr_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.bold = True
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
        # Format Data Rows
        for r_idx, row_data in enumerate(table_rows):
            row_cells = table.rows[r_idx + 1].cells
            bg_color = "F2F2F2" if r_idx % 2 == 1 else "FFFFFF" # Zebra Striping
            for c_idx, cell_text in enumerate(row_data):
                if c_idx < len(row_cells):
                    row_cells[c_idx].text = cell_text.strip()
                    set_cell_background(row_cells[c_idx], bg_color)
                    set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=150, right=150)
                    p = row_cells[c_idx].paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for run in p.runs:
                        run.font.size = Pt(9.0)
                        if "**" in cell_text or "Best" in cell_text or "%" in cell_text:
                            run.font.bold = True

        doc.add_paragraph() # Spacing after table
        in_table = False
        table_headers = []
        table_rows = []

    for line in lines:
        line_str = line.strip()
        
        # Detect Tables
        if line_str.startswith('|') and line_str.endswith('|'):
            parts = [p.strip() for p in line_str.split('|')[1:-1]]
            # Skip delimiter line (| :--- | :--- |)
            if all(set(p).issubset({'-', ':', ' '}) for p in parts):
                continue
            
            if not in_table:
                in_table = True
                table_headers = parts
            else:
                table_rows.append(parts)
            continue
        elif in_table:
            flush_table()

        if line_str.startswith('# '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(line_str[2:])
            run.font.name = 'Calibri'
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)

        elif line_str.startswith('## '):
            p = doc.add_paragraph()
            run = p.add_run(line_str[3:])
            run.font.name = 'Calibri'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)

        elif line_str.startswith('### '):
            p = doc.add_paragraph()
            run = p.add_run(line_str[4:])
            run.font.name = 'Calibri'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)

        elif line_str.startswith('---'):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)

        elif line_str.startswith('1. ') or line_str.startswith('2. ') or line_str.startswith('3. '):
            p = doc.add_paragraph(style='List Number')
            # Handle bold inside item
            text = line_str[3:]
            if '**' in text:
                parts = text.split('**')
                for idx, part in enumerate(parts):
                    run = p.add_run(part)
                    if idx % 2 == 1:
                        run.bold = True
            else:
                p.add_run(text)
            p.paragraph_format.space_after = Pt(3)

        elif line_str:
            p = doc.add_paragraph()
            # Parse inline markdown formatting like **bold**
            if '**' in line_str:
                parts = line_str.split('**')
                for idx, part in enumerate(parts):
                    run = p.add_run(part)
                    if idx % 2 == 1:
                        run.bold = True
            else:
                p.add_run(line_str)
            p.paragraph_format.space_after = Pt(4)

    if in_table:
        flush_table()

    doc.save(docx_path)
    print(f"Successfully converted report to Word Document: '{docx_path}'")

if __name__ == "__main__":
    convert_md_to_docx()
